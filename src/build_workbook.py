"""Build outputs/inventory_policy.xlsx: the same EOQ / safety-stock model as
analysis.py, but with live Excel formulas so assumptions can be changed."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
p = pd.read_csv(OUT / "inventory_policy.csv", dtype={"StockCode": str})
abc = pd.read_csv(OUT / "abc_summary.csv")
acc = pd.read_csv(OUT / "forecast_accuracy.csv")

F = "Arial"
bold = Font(name=F, bold=True); normal = Font(name=F); inp = Font(name=F, color="0000FF")
head = PatternFill("solid", fgColor="DDE7F5")

wb = Workbook()
a = wb.active; a.title = "Assumptions"
a["A1"] = "Inventory policy model: assumptions (blue cells are inputs, change them freely)"
a["A1"].font = bold
rows = [("Ordering cost per purchase order (GBP)", 25, "Assumed"),
        ("Annual holding cost (fraction of unit price)", 0.20, "Assumed"),
        ("Supplier lead time (weeks)", 2, "Assumed"),
        ("Service-level z value (1.645 = 95%)", 1.645, "Standard normal table"),
        ("Weeks per year", 52, ""),
        ("Baseline: orders per year (monthly ordering)", 12, "Comparison case")]
for i, (label, val, note) in enumerate(rows, start=3):
    a.cell(i, 1, label).font = normal
    c = a.cell(i, 2, val); c.font = inp
    a.cell(i, 3, note).font = normal
a["B4"].number_format = "0%"
a["A11"] = ("Demand and price inputs on the Policy sheet come from src/analysis.py "
            "(UCI Online Retail dataset, Dec 2010 to Dec 2011).")
a["A11"].font = normal
a.column_dimensions["A"].width = 48; a.column_dimensions["B"].width = 12
a.column_dimensions["C"].width = 26

s = wb.create_sheet("Policy")
cols = ["Stock code", "Description", "ABC class", "Avg weekly demand (units)",
        "Std dev of weekly forecast error", "Unit price (GBP)", "Annual demand",
        "Holding cost per unit-year", "EOQ (units)", "Safety stock", "Reorder point",
        "Orders per year", "Annual cost at EOQ (GBP)", "Annual cost, monthly ordering (GBP)"]
for j, c in enumerate(cols, 1):
    cell = s.cell(1, j, c); cell.font = bold; cell.fill = head
    cell.alignment = Alignment(wrap_text=True, vertical="top")
A = "Assumptions!"
for i, r in enumerate(p.itertuples(), start=2):
    s.cell(i, 1, r.StockCode); s.cell(i, 2, str(r.Description).title()); s.cell(i, 3, r.Class)
    for j, v in ((4, r.WeeklyDemand), (5, r.SigmaWeekly), (6, r.UnitPrice)):
        s.cell(i, j, round(float(v), 2)).font = inp
    s.cell(i, 7, f"=D{i}*{A}$B$7")
    s.cell(i, 8, f"=F{i}*{A}$B$4")
    s.cell(i, 9, f"=SQRT(2*G{i}*{A}$B$3/H{i})")
    s.cell(i, 10, f"={A}$B$6*E{i}*SQRT({A}$B$5)")
    s.cell(i, 11, f"=D{i}*{A}$B$5+J{i}")
    s.cell(i, 12, f"=G{i}/I{i}")
    s.cell(i, 13, f"=G{i}/I{i}*{A}$B$3+I{i}/2*H{i}")
    s.cell(i, 14, f"={A}$B$8*{A}$B$3+(G{i}/{A}$B$8)/2*H{i}")
    for j in range(1, 15):
        if s.cell(i, j).font != inp: s.cell(i, j).font = normal
        if j >= 4: s.cell(i, j).number_format = "#,##0.00" if j in (6, 8) else "#,##0"
last = len(p) + 1; t = last + 2
s.cell(t, 2, "Total").font = bold
for j in (13, 14):
    L = get_column_letter(j)
    c = s.cell(t, j, f"=SUM({L}2:{L}{last})"); c.font = bold; c.number_format = "#,##0"
s.cell(t + 1, 2, "Cost reduction vs monthly ordering").font = bold
c = s.cell(t + 1, 13, f"=1-M{t}/N{t}"); c.font = bold; c.number_format = "0.0%"
for j, wd in enumerate([11, 38, 8, 13, 15, 10, 11, 12, 10, 10, 10, 10, 14, 16], 1):
    s.column_dimensions[get_column_letter(j)].width = wd
s.freeze_panes = "C2"; s.row_dimensions[1].height = 46

r = wb.create_sheet("Results")
r["A1"] = "ABC classification (all SKUs)"; r["A1"].font = bold
for j, c in enumerate(abc.columns, 1):
    cell = r.cell(2, j, c); cell.font = bold; cell.fill = head
for i, row in enumerate(abc.itertuples(index=False), start=3):
    for j, v in enumerate(row, 1):
        r.cell(i, j, v).font = normal
r["A8"] = "Forecast error on the 8-week hold-out (WAPE, lower is better)"; r["A8"].font = bold
for j, c in enumerate(acc.columns, 1):
    cell = r.cell(9, j, c); cell.font = bold; cell.fill = head
for i, row in enumerate(acc.itertuples(index=False), start=10):
    for j, v in enumerate(row, 1):
        r.cell(i, j, v).font = normal
r["A14"] = "Values on this sheet are written by src/analysis.py."; r["A14"].font = normal
r.column_dimensions["A"].width = 26
for L in "BCDE": r.column_dimensions[L].width = 16

wb.save(OUT / "inventory_policy.xlsx")
print("saved")
